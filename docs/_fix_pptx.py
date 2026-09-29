"""One-off script to apply review fixes to the final submission pptx.
Run once, then delete (not part of the generator pipeline)."""
import sys
sys.path.insert(0, r'd:\Expert_Task\Catchy\docs')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

import generate_pptx as gp

FINAL_PATH = r"d:\Expert_Task\Catchy\docs\[Expert_Task][hoang5.nguyen]Catchy_Android_Log_Viewer_final.pptx"

# (old substring, new substring) - applied at run level, global across all slides
GLOBAL_REPLACEMENTS = [
    ("Keboard Shortcuts", "Keyboard Shortcuts"),
    ("Seach ", "Search "),
    ("Show the hint of the history whey typing the work to search (Ctrl + Space)",
     "Show a history hint when typing the word to search (Ctrl + Space)"),
    ("Write some script(.bat, .py, \u2026) to merged the log fist",
     "Write a script (.bat, .py, etc.) to merge the logs first"),
    ("Can use the whole screen to show a long line of long",
     "Can use the whole screen to show a single long log line"),
    ("Click the bookmark to jump the line is marked and see around source code",
     "Click a bookmark to jump to the marked line and view the surrounding log context"),
    ("Search with two layers: First search by tag, then search by word based on data are searched by tag",
     "Two-layer search: first filter by tag, then search by keyword within the tag-filtered results"),
]


def replace_in_text_frame(tf, replacements):
    for para in tf.paragraphs:
        for run in para.runs:
            for old, new in replacements:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def walk_shapes(shapes, replacements):
    for shape in shapes:
        if shape.has_text_frame:
            replace_in_text_frame(shape.text_frame, replacements)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    replace_in_text_frame(cell.text_frame, replacements)
        if shape.shape_type == 6:  # group
            walk_shapes(shape.shapes, replacements)


prs = Presentation(FINAL_PATH)

# 1. Global typo/wording fixes across all slides
for slide in prs.slides:
    walk_shapes(slide.shapes, GLOBAL_REPLACEMENTS)

# 2. Slide-specific fixes (position-sensitive, must not be applied globally)
slides = list(prs.slides)

# Slide 1 (index 0): "Screen Copy" -> "Scrcpy" to match terminology used elsewhere
walk_shapes(slides[0].shapes, [("Screen Copy", "Scrcpy")])

# Slide 22 (index 21, "Keyboard Shortcuts" after typo fix): fix leftover subtitle
walk_shapes(slides[21].shapes, [("Project Outcome and Value Delivered", "Quick Reference for Power Users")])

# Slides 20 & 21 (index 19, 20): give distinct subtitles instead of repeating slide 19's
walk_shapes(slides[19].shapes, [("Easy-to-use workflow and improved capability", "Layout Customization for Multiple Use Cases")])
walk_shapes(slides[20].shapes, [("Easy-to-use workflow and improved capability", "Dark Theme Support")])

# 3. Add a methodology note under the Performance Evaluation table (slide 4, index 3)
perf_slide = slides[3]
gp.add_text_block(
    perf_slide,
    "Methodology: [fill in test machine spec, e.g. CPU/RAM/disk] - representative Android logcat files, "
    "each size measured end-to-end from file-open to fully rendered table.",
    0.8, 6.5, 11.6, 0.6, 13, gp.MUTED
)

# 4. Add a new Tech Stack & Build slide, inserted right before the Conclusion slide
def make_tech_stack_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gp.set_background(slide)
    gp.add_title(slide, 'Tech Stack & Build', subtitle='Technology, Dependencies, and Packaging')
    bullets = [
        'Language: Python 3.10+',
        'UI framework: PySide6 (Qt 6) with QML for the view layer',
        'Architecture: MVVM-style Controller facade + Worker (QThread) for file I/O and live streaming',
        'Key dependency: pyperclip (clipboard support)',
        'Packaging: PyInstaller -> standalone Catchy.exe, no Python install required',
        'Source layout: components/ (Python backend) - qmls/ (QML views) - styles/ (theming)'
    ]
    gp.add_bullets(slide, bullets, 0.9, 1.9, 11.2, 4.5, font_size=22)
    return slide


make_tech_stack_slide(prs)

# Move the newly appended slide to just before the Conclusion slide (was last, index -1)
sldIdLst = prs.slides._sldIdLst
sld_ids = list(sldIdLst)
new_sld_id = sld_ids[-1]
sldIdLst.remove(new_sld_id)
sldIdLst.insert(len(sld_ids) - 2, new_sld_id)  # before what was the last slide (Conclusion)

prs.save(FINAL_PATH)
print("Saved fixes to:", FINAL_PATH)
print("Total slides now:", len(prs.slides.__iter__.__self__._sldIdLst))
