import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


OUT_PATH = "D:/Expert_Task/Catchy/docs/Catchy_Project_SlideDeck.pptx"

DARK_BG = RGBColor(17, 20, 28)
PANEL_BG = RGBColor(30, 35, 45)
ACCENT_BLUE = RGBColor(97, 218, 251)
ACCENT_GREEN = RGBColor(152, 195, 121)
ACCENT_PURPLE = RGBColor(198, 120, 221)
ACCENT_RED = RGBColor(224, 108, 117)
ACCENT_GOLD = RGBColor(229, 192, 123)
TEXT = RGBColor(240, 240, 240)
MUTED = RGBColor(180, 186, 194)
WHITE = RGBColor(255, 255, 255)


def set_background(slide, color=DARK_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None, y=0.35):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(12.0), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = TEXT

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(y + 0.5), Inches(12.0), Inches(0.4))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.LEFT
        run2 = p2.runs[0]
        run2.font.size = Pt(12)
        run2.font.color.rgb = MUTED


def add_bullets(slide, items, x, y, w, h, font_size=22, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(8)
    tf.margin_bottom = Pt(8)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(font_size)
            r.font.color.rgb = color
            r.font.name = 'Segoe UI'
    return box


def add_text_block(slide, text, x, y, w, h, font_size=20, color=TEXT, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = 'Segoe UI'
    return box


def add_table(slide, rows, cols, x, y, w, h, header_fill=None, row_fill=None):
    table = slide.shapes.add_table(len(rows), cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    table.rows[0].height = Inches(0.4)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(15 if r_idx == 0 else 14)
                    r.font.bold = r_idx == 0
                    r.font.color.rgb = WHITE if r_idx == 0 else TEXT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r_idx == 0 else row_fill or PANEL_BG
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(6)
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
    return table


def add_box(slide, x, y, w, h, text, fill_color, text_color=WHITE, font_size=18, border_color=None, radius=0.12):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill_color
    if border_color:
        box.line.color.rgb = border_color
    else:
        box.line.color.rgb = fill_color
    box.line.width = Pt(1)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = text
    for r in p.runs:
        r.font.size = Pt(font_size)
        r.font.bold = True
        r.font.color.rgb = text_color
        r.font.name = 'Segoe UI'
    return box


def add_arrow(slide, x1, y1, x2, y2, color=ACCENT_BLUE, width=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x1), Inches(y1), Inches(x2 - x1), Inches(y2 - y1))
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def make_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.18))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_BLUE
    accent.line.fill.background()

    heading = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(8.5), Inches(1.2))
    tf = heading.text_frame
    p = tf.paragraphs[0]
    p.text = 'Catchy'
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(32)
        r.font.bold = True
        r.font.color.rgb = ACCENT_BLUE
        r.font.name = 'Segoe UI'

    sub = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(9.2), Inches(0.8))
    tf2 = sub.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = 'Desktop Log Viewer'
    p2.alignment = PP_ALIGN.LEFT
    for r in p2.runs:
        r.font.size = Pt(28)
        r.font.bold = False
        r.font.color.rgb = WHITE
        r.font.name = 'Segoe UI'

    bullet_box = slide.shapes.add_textbox(Inches(0.9), Inches(3.3), Inches(6.5), Inches(2.6))
    tf3 = bullet_box.text_frame
    for idx, text in enumerate([
        'Modern log analysis tool for embedded and Android developers',
        'Python + PySide6 (Qt/QML)',
        'Standalone .exe with no installation required'
    ]):
        p3 = tf3.paragraphs[0] if idx == 0 else tf3.add_paragraph()
        p3.text = '• ' + text
        p3.level = 0
        p3.bullet = True
        for r in p3.runs:
            r.font.size = Pt(21)
            r.font.color.rgb = TEXT
            r.font.name = 'Segoe UI'

    # Minimal visual block on right
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.6), Inches(1.7), Inches(3.8), Inches(3.4))
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL_BG
    panel.line.color.rgb = ACCENT_BLUE
    panel.line.width = Pt(2)
    tf4 = panel.text_frame
    tf4.word_wrap = True
    p4 = tf4.paragraphs[0]
    p4.alignment = PP_ALIGN.CENTER
    p4.text = 'Log file\nLive stream\nSearch\nBookmark\nScrcpy'
    for r in p4.runs:
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = ACCENT_BLUE
        r.font.name = 'Segoe UI'


def make_problem_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Problem Statement', subtitle='Why embedded debugging gets painful')
    rows = [
        ['Problem', 'Impact'],
        ['Log files are huge (100MB–500MB+)', 'Hard to navigate; editors freeze'],
        ['Finding relevant logs is tedious', 'Manual grep / Ctrl+F in text editors'],
        ['Filtering by process/tag requires CLI', 'Steep learning curve and slow workflow'],
        ['Switching between live stream and file analysis', 'Multiple tools needed'],
        ['Losing track of important lines', 'No way to mark and jump back'],
        ['No quick way to share snippets', 'Copy-paste from terminal is messy']
    ]
    add_table(slide, rows, 2, 0.5, 1.5, 12.2, 4.7, header_fill=ACCENT_BLUE, row_fill=PANEL_BG)


def make_features_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Solution — Key Features')
    points = [
        'Log file loading — Open and parse large files with progress indicator',
        'Live logcat streaming — Real-time ADB logcat with auto-scroll',
        'Process filters — Named color-coded filters, persisted to JSON',
        'Multi-keyword search — Highlight multiple terms with different colors',
        'Bookmarks — Save/jump to important log lines',
        'Scrcpy integration — Mirror Android screen alongside log viewing',
        'Column visibility — Show/hide columns for focused reading',
        'Theme support — Dark and Light modes',
        'Standalone executable — Distributable without Python installation'
    ]
    add_bullets(slide, points, 0.7, 1.5, 11.8, 5.0, font_size=18)


def make_architecture_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Architecture Overview')
    add_text_block(slide, 'MVVM architecture', 0.7, 1.2, 2.4, 0.4, 16, MUTED, False)

    left = add_box(slide, 0.9, 2.0, 2.9, 1.3, 'QML Layer\nmain.qml • LogViewTable\nPanels • Styles', ACCENT_BLUE, font_size=15)
    mid = add_box(slide, 4.6, 2.0, 2.8, 1.3, 'Controller\nFacade', ACCENT_GREEN, font_size=16)
    right_top = add_box(slide, 8.0, 1.2, 2.5, 0.9, 'LogModel', ACCENT_PURPLE, font_size=15)
    right_mid = add_box(slide, 8.0, 2.3, 2.5, 0.9, 'FilterLog', ACCENT_PURPLE, font_size=15)
    right_bottom = add_box(slide, 8.0, 3.4, 2.5, 0.9, 'SearchLog', ACCENT_PURPLE, font_size=15)
    lower = add_box(slide, 3.8, 4.5, 2.5, 0.9, 'Bookmark\nHelper', ACCENT_RED, font_size=15)

    summary = add_text_block(slide, 'Pattern: MVVM • Python = ViewModel/Service • QML = View\nBackend: Controller orchestrates all services\nFrontend: Declarative QML UI with data binding', 0.8, 5.3, 10.8, 1.2, 18, TEXT)


def make_large_file_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Large File Loading — How It Works')
    flow = [
        'User opens file',
        'Worker thread starts',
        'Binary read (64KB chunks)',
        'Format pre-detection',
        'Parse lines with one fast parser',
        'Progress emitted every 5,000 lines',
        'Single model update when complete'
    ]
    x = 0.8
    w = 1.55
    y = 1.8
    for idx, text in enumerate(flow):
        box = add_box(slide, x + idx * 1.85, y, w, 0.8, text, ACCENT_BLUE if idx % 2 == 0 else ACCENT_GREEN, font_size=11)
    add_text_block(slide, 'Key optimizations', 0.8, 3.2, 2.6, 0.4, 18, MUTED)
    rows = [
        ['Technique', 'Benefit'],
        ['64KB binary buffer', 'Faster than Python readline()'],
        ['Format pre-detection', '1 parser/line instead of trying 3'],
        ['Tag string interning', '~60% heap reduction for duplicates'],
        ['Single data store', 'No duplicate _logDict — 1x memory'],
        ['Worker thread', 'UI stays responsive during load']
    ]
    add_table(slide, rows, 2, 0.7, 3.6, 11.8, 2.8, header_fill=ACCENT_RED, row_fill=PANEL_BG)


def make_streaming_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Live Logcat Streaming — How It Works')
    add_box(slide, 0.8, 1.8, 2.0, 1.0, 'ADB logcat\n(stdout)', ACCENT_RED, font_size=16)
    add_box(slide, 3.1, 1.8, 2.1, 1.0, 'Writer thread\n50MB rotation', ACCENT_GOLD, font_size=16)
    add_box(slide, 5.7, 1.8, 2.0, 1.0, '.log files\n(on disk)', ACCENT_BLUE, font_size=16)
    add_box(slide, 8.3, 1.8, 2.0, 1.0, 'Reader thread\nparse + buffer', ACCENT_GREEN, font_size=16)
    add_box(slide, 10.9, 1.8, 1.7, 1.0, 'Main thread\nBatch insert', ACCENT_PURPLE, font_size=15)
    add_text_block(slide, 'Rate limiting: flush timer at 100ms → UI updates max 10×/sec\nFile rotation: split at 50MB to prevent unbounded growth\nAuto-scroll pause: stops UI flushing for smooth manual scrolling', 0.8, 3.4, 11.0, 1.8, 19, TEXT)


def make_threading_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Threading Model')
    rows = [
        ['Thread', 'Purpose'],
        ['Main (UI)', 'QML rendering and property updates'],
        ['_loadLogFileThread', 'File loading and saving'],
        ['_logcatThread', 'ADB → file writer (50MB rotation)'],
        ['_logcatStreamThread', 'File → buffer reader'],
        ['_filterThread', 'Filter add/update operations'],
        ['_searchThread', 'Search execution']
    ]
    add_table(slide, rows, 2, 0.8, 1.7, 11.8, 3.7, header_fill=ACCENT_BLUE, row_fill=PANEL_BG)
    add_text_block(slide, 'All workers use Worker(QObject) with taskCompleted signal — clean thread lifecycle management.', 0.8, 5.8, 11.0, 0.8, 17, TEXT)


def make_filter_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Filter & Color System')
    bullets = [
        'Lazy computation: colors are NOT stored per row — computed on-demand in data() role',
        'Filter colors: regex match against process name → assigned color',
        'Level colors: direct dict lookup (V / D / I / W / E / F)',
        'Proxy model: custom Python proxy with precomputed index list',
        ' _rebuild(): O(N) single pass + beginResetModel/endResetModel',
        'Streaming support: new rows checked and appended incrementally'
    ]
    add_bullets(slide, bullets, 0.8, 1.7, 11.8, 4.5, font_size=18)


def make_performance_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Performance Results')
    rows = [
        ['Metric', 'Implementation'],
        ['File loading', 'Worker thread + 64KB binary chunks — UI responsive'],
        ['Memory', 'Single store + tag interning — no duplicates'],
        ['Parsing', 'Format pre-detection — 1 parser per line'],
        ['Filter change', 'Lazy color in data() — no full recolor'],
        ['Save', 'Chunked 10K rows with progress — bounded memory'],
        ['Streaming', '100ms batch flush — decoupled from parse rate']
    ]
    add_table(slide, rows, 2, 0.8, 1.6, 11.8, 4.2, header_fill=ACCENT_GREEN, row_fill=PANEL_BG)
    add_text_block(slide, 'Target: Handle 100MB–500MB+ files smoothly', 0.8, 6.0, 8.0, 0.6, 20, ACCENT_GREEN, True)


def make_benefits_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Benefits')
    bullets = [
        'Productivity — No more manual grep or Ctrl+F in text editors',
        'Visual clarity — Color-coded processes and search terms make patterns easy to spot',
        'All-in-one — File viewer + live stream + device mirror in one tool',
        'Cross-workflow — Post-mortem analysis and real-time debugging',
        'Lightweight — Single executable, no heavy IDE required',
        'Customizable — Filters and configs saved per user, ready for the next session'
    ]
    add_bullets(slide, bullets, 0.9, 1.7, 11.6, 4.8, font_size=20)


def make_demo_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Demo Flow')
    steps = [
        'Open a log file -> Show loading progress and table rendering',
        'Add a filter -> Show color-coded process highlighting',
        'Search keywords -> Show multi-color highlights',
        'Bookmark a line -> Show bookmark panel and jump-to',
        'Switch to logcat -> Show live streaming with auto-scroll',
        'Toggle theme -> Show Dark/Light switch',
        'Optional: Scrcpy -> Show device screen mirroring'
    ]
    add_bullets(slide, steps, 1.0, 1.9, 11.2, 4.8, font_size=21)


def make_qa_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Q&A')
    add_text_block(slide, 'Thank you!', 0.8, 1.6, 4.0, 0.6, 28, ACCENT_BLUE, True)
    bullets = [
        'Tech stack: Python 3.10+ / PySide6 / QML',
        'Dependencies: PySide6, pyperclip',
        'Build: PyInstaller → standalone .exe',
        'Source structure: MVVM with Controller facade pattern'
    ]
    add_bullets(slide, bullets, 1.0, 2.4, 10.5, 3.2, font_size=22)


def make_compare_overview_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Catchy vs Bonnibel.exe', subtitle='Positioning and value proposition')
    add_text_block(slide, 'Bonnibel.exe is a useful reference point but it is still limited for modern debugging workflows.', 0.8, 1.6, 8.5, 0.7, 20, TEXT)
    add_text_block(slide, 'Catchy focuses on fast filtering, large log navigation, live device debugging, and a streamlined workflow for engineers.', 0.8, 2.4, 8.8, 0.9, 20, TEXT)
    add_box(slide, 0.9, 3.5, 3.4, 1.6, 'Bonnibel.exe\nLimited workflow\nMostly single-task usage', ACCENT_RED, font_size=18)
    add_box(slide, 4.7, 3.5, 3.7, 1.6, 'Catchy\nAll-in-one log viewer\nLive + file + filter + bookmark', ACCENT_GREEN, font_size=18)
    add_text_block(slide, 'Core idea: reduce context switching and speed up diagnosis for embedded/Android engineers.', 0.8, 5.8, 11.2, 0.8, 19, ACCENT_BLUE)


def make_compare_matrix_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Feature Comparison')
    rows = [
        ['Capability', 'Bonnibel.exe', 'Catchy'],
        ['Large log file handling', 'Basic / limited handling', 'Worker-thread loading with progress'],
        ['Live logcat', 'Limited or separate workflow', 'Integrated live ADB streaming'],
        ['Process filtering', 'Manual, narrow scope', 'Named color-coded filters saved to JSON'],
        ['Search', 'Simple search', 'Multi-keyword highlighting with colors'],
        ['Bookmarks', 'Weak or absent', 'Save/jump to important lines'],
        ['Android screen mirroring', 'Not in focus', 'Scrcpy integration included'],
        ['Theme / UX', 'Basic UI', 'Dark and Light modes'],
        ['Export/share workflow', 'Minimal', 'Clipboard-ready snippets + quick navigation']
    ]
    add_table(slide, rows, 3, 0.5, 1.6, 12.3, 4.9, header_fill=ACCENT_BLUE, row_fill=PANEL_BG)


def make_compare_advantages_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Why Catchy Wins')
    bullets = [
        'Better for large debugging sessions with heavy log data',
        'Faster triage through filters, bookmarks, and color-coded search',
        'Live logcat + file analysis in one workflow',
        'Less context switching for Android and embedded developers',
        'Standalone desktop app with a cleaner user experience',
        'Built to support iterative debugging instead of a limited one-off view'
    ]
    add_bullets(slide, bullets, 1.0, 1.75, 10.8, 4.8, font_size=20)
    add_text_block(slide, 'Outcome: faster diagnosis, cleaner workflow, higher productivity.', 0.9, 6.5, 8.8, 0.6, 19, ACCENT_GREEN, True)


def make_intro_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Catchy', subtitle='Desktop Log Viewer for Embedded and Android Debugging')
    bullets = [
        'A unified platform for log analysis and real-time logcat monitoring',
        'Designed for large-scale log files, structured filtering, and workflow optimization',
        'Improves debugging speed, traceability, and usability compared with manual inspection'
    ]
    add_bullets(slide, bullets, 0.9, 1.9, 6.4, 2.8, font_size=22)
    add_box(slide, 8.3, 1.9, 4.0, 3.4, 'Catchy\nLog viewer\nLive stream\nSearch\nBookmark', ACCENT_BLUE, font_size=20)


def make_bonnibel_problem_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Problem Statement', subtitle='Limitations of the Existing Tool')
    bullets = [
        'Large log files are difficult to process efficiently without noticeable delays',
        'Search and filtering operations are insufficient for rapid debugging workflows',
        'The user must switch contexts between viewing logs and monitoring device activity',
        'Important events are harder to trace and revisit during repeated debugging cycles',
        'The existing workflow does not support smooth live investigation and structured analysis'
    ]
    add_bullets(slide, bullets, 0.9, 1.7, 7.2, 4.0, font_size=21)
    add_box(slide, 8.8, 2.0, 3.5, 2.2, 'Bonnibel.exe\nLimited workflow\nSlower triage', ACCENT_RED, font_size=19)


def make_compare_proof_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Comparative Evidence', subtitle='Legacy Tool Versus Proposed System')

    bonnibel_img = r'D:\Expert_Task\Catchy\imgs\bonnibel\Main_UI.png'
    catchy_img = r'D:\Expert_Task\Catchy\imgs\catchy\Main_UI.png'

    if os.path.exists(bonnibel_img):
        slide.shapes.add_picture(bonnibel_img, Inches(0.8), Inches(1.8), Inches(5.4), Inches(3.6))
    else:
        add_box(slide, 0.9, 1.8, 5.2, 4.4, 'Bonnibel.exe\nLegacy UI', ACCENT_RED, font_size=18)

    if os.path.exists(catchy_img):
        slide.shapes.add_picture(catchy_img, Inches(7.0), Inches(1.8), Inches(5.4), Inches(3.6))
    else:
        add_box(slide, 7.0, 1.8, 5.2, 4.4, 'Catchy\nModern UI', ACCENT_GREEN, font_size=18)

    add_text_block(slide, 'The updated interface demonstrates a clearer layout, reduced operational friction, and a more efficient debugging workflow.', 0.9, 6.0, 10.8, 0.7, 19, ACCENT_BLUE, True)


def make_architecture_story_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'System Architecture')
    add_box(slide, 0.8, 1.8, 2.8, 1.4, 'QML UI', ACCENT_BLUE, font_size=17)
    add_box(slide, 3.9, 1.8, 2.8, 1.4, 'Controller', ACCENT_GREEN, font_size=17)
    add_box(slide, 7.0, 1.8, 2.8, 1.4, 'Log Model', ACCENT_PURPLE, font_size=17)
    add_box(slide, 10.1, 1.8, 2.2, 1.4, 'Filters\nSearch\nBookmark', ACCENT_RED, font_size=16)
    add_text_block(slide, 'The application follows an MVVM-inspired structure with worker threads to support large-file loading and live streaming.', 0.9, 3.8, 10.5, 0.8, 20, TEXT)
    add_text_block(slide, 'The controller coordinates the core logic, while the QML layer ensures responsive visual updates and efficient user interaction.', 0.9, 4.8, 10.8, 0.8, 18, MUTED)


def make_features_story_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Core Functional Features')
    bullets = [
        'Large-file processing with progress indicators and worker-thread execution',
        'Live logcat streaming with auto-scroll and real-time refresh',
        'Multi-keyword highlighting using color-coded matching for rapid analysis',
        'Named filters for process and tag-based debugging workflows',
        'Bookmark system for saving and revisiting critical log entries',
        'Scrcpy integration for device-side visual verification and validation'
    ]
    add_bullets(slide, bullets, 0.9, 1.8, 11.1, 4.7, font_size=20)


def make_performance_update_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Performance Evaluation', subtitle='Measured Loading-Time Comparison')
    rows = [
        ['File size', 'Catchy', 'Bonnibel.exe'],
        ['20 MB', 'under 2s', 'under 2s'],
        ['100 MB', '3s', '9s'],
        ['250 MB', '6s', '20s'],
        ['500 MB', '12s', '45s / occasional UI freeze']
    ]
    add_table(slide, rows, 3, 0.8, 1.8, 11.8, 3.6, header_fill=ACCENT_GREEN, row_fill=PANEL_BG)
    add_text_block(slide, 'Result: Catchy maintains responsiveness even with large log files, whereas the legacy tool experiences noticeable delays.', 0.9, 5.8, 10.8, 0.7, 20, ACCENT_GREEN, True)


def make_evidence_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Evidence of Improvement', subtitle='Major Bottleneck Reduction and Enhanced Usability')
    bullets = [
        '500MB logs load in approximately 12 seconds in Catchy, compared with about 45 seconds and occasional UI freezing in Bonnibel',
        'The adaptive interface maintains responsiveness while streaming live logcat data',
        'Automatic ADB state detection reduces user errors and improves operational reliability',
        'Keyboard shortcuts and workflow automation significantly reduce manual navigation and debugging effort',
        'The system now supports continuous monitoring with a 500MB retention window and smooth scrolling performance'
    ]
    add_bullets(slide, bullets, 0.9, 1.7, 11.2, 4.5, font_size=19)


def make_usability_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'User Experience and Workflow Optimization', subtitle='Automation and Shortcut-Based Efficiency')
    bullets = [
        'Automatically enables Start Stream Logcat and Start Screen Copy when an ADB device is connected',
        'Automatically disables the same actions when no compatible device is available',
        'Keyboard shortcuts accelerate the debugging workflow: Ctrl+F search, Ctrl+B columns, Ctrl+T filter edit, Ctrl+M bookmark panel, Ctrl+P auto-scroll',
        'Smooth scrolling, drag-to-open file handling, and an efficient side-panel layout reduce friction during analysis',
        'A 500MB retention window preserves recent activity and supports continuous live monitoring without losing context'
    ]
    add_bullets(slide, bullets, 0.9, 1.7, 11.2, 4.8, font_size=20)


def make_test_story_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Testing & Validation')
    rows = [
        ['Validation area', 'Result'],
        ['Large log loading', 'Responsive while reading and parsing huge files'],
        ['Live stream', 'Stable updates and auto-scroll control'],
        ['Filters', 'Fast matching and color grouping'],
        ['Search', 'Highlight multiple keywords correctly'],
        ['Bookmarks', 'Quick jump and persistence'],
        ['UI behavior', 'Dark/light theme and table responsiveness'],
        ['ADB device handling', 'Auto enable/disable based on connection state']
    ]
    add_table(slide, rows, 2, 0.8, 1.7, 11.7, 4.2, header_fill=ACCENT_BLUE, row_fill=PANEL_BG)


def make_conclusion_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Conclusion', subtitle='Project Outcome and Value Delivered')
    bullets = [
        'Catchy is a practical and user-centered log analysis tool for embedded and Android debugging workflows',
        'It resolves key limitations in large-file handling, filtering efficiency, and repetitive manual operations',
        'The solution improves real-world productivity through performance optimization, automation, and streamlined interaction design',
        'The final result offers a cleaner, faster, and more reliable debugging workflow than the reference tool'
    ]
    add_bullets(slide, bullets, 1.0, 1.8, 10.5, 4.1, font_size=22)
    add_text_block(slide, 'Expert project outcome: a validated, evidence-backed solution designed for practical engineering use.', 1.0, 6.0, 10.0, 0.7, 19, ACCENT_GREEN, True)


def add_feature_image_slide(prs, title, subtitle, image_path, caption, x=0.8, y=1.6, w=5.8, h=4.0):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle=subtitle)
    if os.path.exists(image_path):
        slide.shapes.add_picture(image_path, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        add_box(slide, x, y, w, h, 'Image not available', ACCENT_RED, font_size=18)
    add_text_block(slide, caption, 0.9, 6.1, 11.5, 0.8, 19, TEXT)
    return slide


def make_feature_image_slides(prs):
    feature_specs = [
        (
            'Full Log Viewer UI',
            'Main view with log table and filtering controls',
            r'D:\Expert_Task\Catchy\imgs\catchy\Full_Log_Viewer_UI.png',
            'The full log viewer provides a readable, structured interface for large logs, making analysis more efficient than a text-heavy legacy tool.'
        ),
        (
            'Long Search Box',
            'Quick keyword search and result navigation',
            r'D:\Expert_Task\Catchy\imgs\catchy\TopBar_Long_Search_text_box.png',
            'A dedicated search bar enables fast keyword search, better visibility, and quick navigation to relevant log entries.'
        ),
        (
            'Tag Filter and Bookmark Panel',
            'Filter management and bookmark workflow',
            r'D:\Expert_Task\Catchy\imgs\catchy\Left_Bar_Book_mark.png',
            'The left-side panels allow users to monitor active filters and bookmarks without leaving the main debugging context.'
        ),
        (
            'Tag Filter Panel',
            'Color-coded tag and process filtering',
            r'D:\Expert_Task\Catchy\imgs\catchy\Left_Bar_Tag_Filter.png',
            'Named filters and color coding make it easy to isolate important processes and reduce noise in large log sets.'
        ),
        (
            'Split View for Long Content',
            'Readable handling of expanded log content',
            r'D:\Expert_Task\Catchy\imgs\catchy\Split_UI_for_long_content.png',
            'The split layout keeps long line content readable and avoids the cramped interface that makes debugging difficult.'
        ),
        (
            'Auto-enable Scrcpy',
            'Smart device integration',
            r'D:\Expert_Task\Catchy\imgs\catchy\Auto_Enable_Scrcpy.png',
            'When the device is connected, the app automatically enables the screen-copy workflow for a smoother live debugging experience.'
        )
    ]
    for title, subtitle, image_path, caption in feature_specs:
        add_feature_image_slide(prs, title, subtitle, image_path, caption)


def make_live_proof_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, 'Proof of the New UI and Features', subtitle='Easy-to-use workflow and improved capability')

    catchy_main = r'D:\Expert_Task\Catchy\imgs\catchy\Main_UI.png'
    catchy_features = r'D:\Expert_Task\Catchy\imgs\catchy\Full_Log_Viewer_UI.png'

    if os.path.exists(catchy_main):
        slide.shapes.add_picture(catchy_main, Inches(0.8), Inches(1.8), Inches(5.4), Inches(3.7))
    else:
        add_box(slide, 0.9, 1.8, 5.2, 4.4, 'Catchy UI', ACCENT_GREEN, font_size=20)

    if os.path.exists(catchy_features):
        slide.shapes.add_picture(catchy_features, Inches(7.0), Inches(1.8), Inches(5.4), Inches(3.7))
    else:
        add_box(slide, 7.0, 1.8, 5.2, 4.4, 'Catchy Features', ACCENT_BLUE, font_size=20)

    add_text_block(slide, 'The updated interface and feature set show a cleaner, more usable, and more efficient debugging workflow.', 0.9, 6.0, 10.8, 0.7, 18, TEXT)


def generate_expert_summary_md():
    summary = """# Expert Project Proof Summary

## Project: Catchy

### 1. Problem addressed
The original comparison target, Bonnibel.exe, shows major limitations in large log handling and workflow usability for Android and embedded debugging.

### 2. Performance evidence
- 20MB: under 2s for both tools
- 100MB: Catchy ~3s vs Bonnibel ~9s
- 250MB: Catchy ~6s vs Bonnibel ~20s
- 500MB: Catchy ~12s vs Bonnibel ~45s and occasional UI freeze

This demonstrates a clear speed advantage and better stability for large log files.

### 3. Usability evidence
- Auto-enable Start Stream Logcat and Start Screen Copy when an ADB device is connected
- Auto-disable those controls when the device is missing
- Shortcut support for faster work: Ctrl+F, Ctrl+B, Ctrl+T, Ctrl+M, Ctrl+P
- Smooth scrolling, drag-to-open file, and filtered browsing improve daily debugging productivity

### 4. Technical architecture
Catchy is implemented with Python + PySide6 + QML using an MVVM-like architecture. The main structure is:
- QML UI layer
- Controller facade
- Log model and proxy model
- Filter management, search, bookmarks, and helper services
- Worker threads for long-running I/O and live streaming

### 5. User-facing validation
The project includes evidence from actual screenshots showing the cleaner interface, richer feature set, and easier workflow compared with the legacy tool.

### 6. Conclusion
Catchy was developed to provide a faster, clearer, and more practical debugging experience for large log analysis and Android device workflows.
"""
    out_path = r'D:\Expert_Task\Catchy\docs\Expert_Project_Proof.md'
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(summary)
    return out_path


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    make_intro_slide(prs)
    make_bonnibel_problem_slide(prs)
    make_compare_proof_slide(prs)
    make_performance_update_slide(prs)
    make_evidence_slide(prs)
    make_architecture_story_slide(prs)
    make_features_story_slide(prs)
    make_feature_image_slides(prs)
    make_usability_slide(prs)
    make_test_story_slide(prs)
    make_live_proof_slide(prs)
    make_conclusion_slide(prs)

    prs.save(OUT_PATH)
    summary_path = generate_expert_summary_md()
    print(f'Presentation created: {OUT_PATH}')
    print(f'Project proof summary created: {summary_path}')


if __name__ == '__main__':
    build_presentation()
